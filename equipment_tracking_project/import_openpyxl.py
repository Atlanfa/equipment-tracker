import openpyxl
from equipment_tracking_app.models import Equipment

def process_xlsx_file(file_path):
    # Load the workbook
    workbook = openpyxl.load_workbook(file_path)

    # Iterate through each sheet in the workbook
    for sheet_name in workbook.sheetnames:
        sheet = workbook[sheet_name]
        print(f"Processing sheet: {sheet_name}")

        # Get the first row (subcategory names)
        subcategories = [cell.value for cell in sheet[1]]

        # Iterate through the columns below the first row
        for col_idx, subcategory in enumerate(subcategories, start=1):
            if subcategory:  # Ensure the subcategory name is not None
                for row in sheet.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                    for cell in row:
                        equipment_name = cell.value
                        if equipment_name:  # Ensure the cell is not empty
                            # Save the equipment name to the database
                            equipment, created = Equipment.objects.get_or_create(
                                name=equipment_name.strip()
                            )
                            if created:
                                print(f"Added new equipment: {equipment_name}")
                            else:
                                print(f"Equipment already exists: {equipment_name}")