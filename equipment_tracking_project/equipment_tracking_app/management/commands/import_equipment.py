from django.core.management.base import BaseCommand
from equipment_tracking_app.models import Equipment
from openpyxl import load_workbook

class Command(BaseCommand):
    help = "Import equipment names from an .xlsx file"

    def add_arguments(self, parser):
        parser.add_argument('file_path', type=str, help="Path to the .xlsx file")

    def handle(self, *args, **kwargs):
        file_path = kwargs['file_path']
        workbook = load_workbook(file_path)

        for sheet_name in workbook.sheetnames:
            sheet = workbook[sheet_name]
            self.stdout.write(f"Processing sheet: {sheet_name}")

            subcategories = [cell.value for cell in sheet[1]]

            for col_idx, subcategory in enumerate(subcategories, start=1):
                if subcategory:
                    for row in sheet.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
                        for cell in row:
                            equipment_name = cell.value
                            if equipment_name:
                                equipment, created = Equipment.objects.get_or_create(
                                    name=equipment_name.strip()
                                )
                                if created:
                                    self.stdout.write(f"Added new equipment: {equipment_name}")
                                else:
                                    self.stdout.write(f"Equipment already exists: {equipment_name}")